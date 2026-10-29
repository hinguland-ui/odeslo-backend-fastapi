"""The super admin report: filters change the numbers, exports match, other roles are refused."""

import io
import os
from datetime import date, timedelta

from openpyxl import load_workbook

from tests.conftest import API, auth


def _get(client, token, path, **params):
    return client.get(f"{API}/reports/{path}", headers=auth(token), params=params)


def test_overview_filters_and_access(client, state):
    sa, soa, r1 = state["sa"], state["soa"], state["r1"]
    today = date.today()

    # only the super admin sees platform reports
    assert _get(client, soa, "overview").status_code == 403
    assert _get(client, r1, "overview").status_code == 403

    full = _get(client, sa, "overview", date_from=str(today - timedelta(days=60)), date_to=str(today)).json()
    k = full["kpis"]
    assert k["gross_revenue"] > 0 and k["payments"] > 0 and k["societies"] >= 2
    assert k["gross_revenue"] == round(k["booking_revenue"] + k["maintenance_revenue"], 2)
    assert k["net_revenue"] == round(k["gross_revenue"] - k["refunded"], 2)
    assert len(full["societies"]) == k["societies"]
    # the time series covers every day of the range, zero days included
    assert len(full["timeseries"]) == 61 and full["filters"]["group_by"] == "day"
    assert round(sum(t["booking"] + t["maintenance"] for t in full["timeseries"]), 2) == k["gross_revenue"]
    assert round(sum(s["total_revenue"] for s in full["societies"]), 2) == k["gross_revenue"]
    assert round(sum(p["amount"] for p in full["payments"]), 2) == k["gross_revenue"]

    # society filter narrows everything to that society
    sid = state["society_id"]
    one = _get(client, sa, "overview", date_from=str(today - timedelta(days=60)), date_to=str(today), society_ids=str(sid)).json()
    assert [s["society_id"] for s in one["societies"]] == [sid] and one["filters"]["societies"] != ["All societies"]
    assert one["kpis"]["gross_revenue"] <= k["gross_revenue"]
    assert all(p["society"] == one["societies"][0]["society"] for p in one["payments"])

    # payment mode, category and a date range with nothing in it
    online = _get(client, sa, "overview", date_from=str(today - timedelta(days=60)), date_to=str(today), mode="online").json()
    assert all(p["mode"] == "online" for p in online["payments"])
    cat = _get(client, sa, "overview", date_from=str(today - timedelta(days=60)), date_to=str(today), category="ac").json()
    assert all(p["type"] == "booking" for p in cat["payments"])
    empty = _get(client, sa, "overview", date_from="2020-01-01", date_to="2020-01-31").json()
    assert empty["kpis"]["gross_revenue"] == 0 and empty["payments"] == [] and len(empty["timeseries"]) == 31

    # weekly / monthly grouping and bad input
    month = _get(client, sa, "overview", date_from=str(today - timedelta(days=200)), date_to=str(today), group_by="month").json()
    assert month["filters"]["group_by"] == "month" and 6 <= len(month["timeseries"]) <= 8
    assert _get(client, sa, "overview", date_from="2026-02-01", date_to="2026-01-01").status_code == 400
    assert _get(client, sa, "overview", date_from="2000-01-01", date_to="2026-01-01").status_code == 400


def test_exports_match_the_filters(client, state):
    sa, soa = state["sa"], state["soa"]
    today = date.today()
    q = {"date_from": str(today - timedelta(days=60)), "date_to": str(today)}
    data = _get(client, sa, "overview", **q).json()

    assert _get(client, soa, "export", format="xlsx", **q).status_code == 403
    assert _get(client, sa, "export", format="csv", **q).status_code == 422

    x = _get(client, sa, "export", format="xlsx", **q)
    assert x.status_code == 200 and "spreadsheetml" in x.headers["content-type"]
    assert ".xlsx" in x.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(x.content))
    assert {"Summary", "Revenue", "Payments", "Societies", "Bookings & deals", "Maintenance", "Complaints"} <= set(wb.sheetnames)
    pay = wb["Payments"]
    assert pay["A1"].value == "Paid on" and pay["H1"].value == "Amount"
    assert pay.max_row - 2 == len(data["payments"])  # header + rows + the total row
    assert pay["A2"].number_format == "dd-mmm-yyyy hh:mm" and pay["H2"].number_format.startswith('"₹"')
    rev = wb["Revenue"]
    assert rev.max_row - 2 == len(data["timeseries"])  # one row per period (+ header, + total)
    soc = wb["Societies"]
    assert soc.max_row - 2 == len(data["societies"])

    # the PDF is a real, multi-section document
    p = _get(client, sa, "export", format="pdf", **q)
    assert p.status_code == 200 and p.content.startswith(b"%PDF") and len(p.content) > 8000
    assert "Odeslo-report-" in p.headers["content-disposition"]

    # a filtered export only has the filtered society
    sid = state["society_id"]
    xf = _get(client, sa, "export", format="xlsx", society_ids=str(sid), **q)
    wbf = load_workbook(io.BytesIO(xf.content))
    assert wbf["Societies"].max_row - 2 == 1
    assert sid and "Societies" in wbf.sheetnames

    keep = os.environ.get("REPORT_SAMPLE_DIR")
    if keep:  # a hand-inspection hook: REPORT_SAMPLE_DIR=/some/dir pytest ...
        open(os.path.join(keep, "report.xlsx"), "wb").write(x.content)
        open(os.path.join(keep, "report.pdf"), "wb").write(p.content)
