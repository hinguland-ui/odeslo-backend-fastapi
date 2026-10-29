"""Tampering with a request must not change what is charged, confirmed or refunded.

The server never takes an amount from the client: it works it out from its own records, and a payment is confirmed only
by the gateway's signature / its own report of the amount paid.
"""

import json
from datetime import date, datetime, timedelta

from tests.conftest import API, auth


def _deal(client, soa, price=700, max_units=10):
    cutoff = (datetime.now() + timedelta(days=2)).isoformat(timespec="minutes")
    r = client.post(f"{API}/group-deals", headers=auth(soa), json={
        "title": "Tamper test", "category": "ac", "service_date": str(date.today() + timedelta(days=6)),
        "slots": ["09:00-12:00"], "price_per_unit": price, "max_units": max_units, "cutoff_date": cutoff})
    assert r.status_code == 201, r.text
    deal_id = r.json()["id"]
    assert client.post(f"{API}/group-deals/{deal_id}/publish", headers=auth(soa)).status_code == 200
    return deal_id


def _verify(client, token, order, signature="mock_signature", pay_id="pay_x"):
    return client.post(f"{API}/payments/verify", headers=auth(token), json={
        "payment_id": order["payment_id"], "gateway_order_id": order["gateway_order_id"],
        "gateway_payment_id": pay_id, "signature": signature})


def test_client_cannot_choose_the_price(client, state):
    soa, r1 = state["soa"], state["r1"]
    deal_id = _deal(client, soa, price=700)

    # extra "amount" / "price" fields in the body are ignored: the booking costs what the deal says
    b = client.post(f"{API}/bookings", headers=auth(r1), json={
        "deal_id": deal_id, "units_count": 2, "slot": "09:00-12:00", "amount": 1, "price_per_unit": 1, "total_amount": 0})
    assert b.status_code == 201, b.text
    booking = b.json()
    assert booking["amount"] == 1400

    # same for the payment order: only "what to pay for" is sent, the amount comes from the server
    order = client.post(f"{API}/payments/initiate", headers=auth(r1), json={
        "ref_type": "booking", "ref_id": booking["id"], "amount": 1, "amount_paise": 100}).json()
    assert order["amount"] == 1400 and order["amount_paise"] == 140000
    state["tamper"] = {"deal": deal_id, "booking": booking["id"], "order": order}


def test_forged_confirmation_is_refused(client, state):
    r1, r2, order = state["r1"], state["r2"], state["tamper"]["order"]

    # somebody else cannot confirm (or even see) this payment
    assert _verify(client, r2, order).status_code == 404
    # a made-up signature does not confirm it
    bad = _verify(client, r1, order, signature="forged")
    assert bad.status_code == 400
    # a signature for a different order id does not either
    other = {**order, "gateway_order_id": "mock_order_other"}
    assert _verify(client, r1, other).status_code == 400
    booking = client.get(f"{API}/bookings/{state['tamper']['booking']}", headers=auth(r1)).json()
    assert booking["payment_status"] == "unpaid" and booking["status"] == "pending_payment"


def test_webhook_amount_must_match(client, state, monkeypatch):
    from app.integrations.payment_gateway import RazorpayGateway

    # stand-in for the HMAC check with the webhook secret: only the signature "good" is genuine
    monkeypatch.setattr(RazorpayGateway, "verify_webhook", lambda self, body, sig: sig == "good")
    soa, r1 = state["soa"], state["r1"]
    booking_id = state["tamper"]["booking"]
    order = client.post(f"{API}/payments/initiate", headers=auth(r1), json={"ref_type": "booking", "ref_id": booking_id}).json()

    def hook(amount_paise, signature="good"):
        body = {"event": "payment.captured", "payload": {"payment": {"entity": {
            "id": "pay_hook", "order_id": order["gateway_order_id"], "amount": amount_paise, "status": "captured"}}}}
        return client.post(f"{API}/payments/webhook/razorpay", content=json.dumps(body),
                           headers={"X-Razorpay-Signature": signature, "content-type": "application/json"})

    assert hook(140000, signature="wrong").status_code == 403  # unsigned / forged webhooks are refused
    assert hook(100).status_code == 200  # signed, but the gateway reports Rs 1 instead of Rs 1400
    assert client.get(f"{API}/payments/{order['payment_id']}", headers=auth(r1)).json()["status"] != "success"
    booking = client.get(f"{API}/bookings/{booking_id}", headers=auth(r1)).json()
    assert booking["payment_status"] == "unpaid"
    admin_notes = client.get(f"{API}/notifications", headers=auth(soa)).json()
    assert any("mismatch" in n["title"].lower() for n in admin_notes["items"])

    assert hook(140000).status_code == 200  # the right amount confirms it
    booking = client.get(f"{API}/bookings/{booking_id}", headers=auth(r1)).json()
    assert booking["payment_status"] == "paid" and booking["status"] == "confirmed"
    state["tamper"]["paid_order"] = order


def test_second_payment_for_a_paid_booking_is_refunded(client, state):
    soa, r1 = state["soa"], state["r1"]
    # a second checkout opened before the first was paid (a new booking, paid twice)
    deal_id = _deal(client, soa, price=300)
    booking = client.post(f"{API}/bookings", headers=auth(r1), json={"deal_id": deal_id, "units_count": 1, "slot": "09:00-12:00"}).json()
    first = client.post(f"{API}/payments/initiate", headers=auth(r1), json={"ref_type": "booking", "ref_id": booking["id"]}).json()
    second = client.post(f"{API}/payments/initiate", headers=auth(r1), json={"ref_type": "booking", "ref_id": booking["id"]}).json()
    assert _verify(client, r1, first, pay_id="pay_a").status_code == 200
    assert _verify(client, r1, second, pay_id="pay_b").status_code == 200

    b = client.get(f"{API}/bookings/{booking['id']}", headers=auth(r1)).json()
    assert b["payment_status"] == "paid" and b["status"] == "confirmed"  # credited once
    refunds = client.get(f"{API}/refunds?status=requested", headers=auth(soa)).json()["items"]
    assert [r["amount"] for r in refunds if r["payment_id"] == second["payment_id"]] == [300.0]  # the extra one goes back
    # and a paid booking cannot be paid again at all
    assert client.post(f"{API}/payments/initiate", headers=auth(r1), json={"ref_type": "booking", "ref_id": booking["id"]}).status_code == 400


def test_mock_gateway_never_runs_in_production(client, state, monkeypatch):
    from app.core.config import settings

    r1 = state["r1"]
    deal_id = _deal(client, state["soa"], price=100)
    booking = client.post(f"{API}/bookings", headers=auth(r1), json={"deal_id": deal_id, "units_count": 1, "slot": "09:00-12:00"}).json()
    order = client.post(f"{API}/payments/initiate", headers=auth(r1), json={"ref_type": "booking", "ref_id": booking["id"]}).json()

    monkeypatch.setattr(settings, "ENVIRONMENT", "production")
    # with the test gateway ("mock_signature" always passes) live, anyone could mark themselves paid
    assert client.post(f"{API}/payments/initiate", headers=auth(r1), json={"ref_type": "booking", "ref_id": booking["id"]}).status_code == 400
    assert _verify(client, r1, order).status_code == 400
    assert client.get(f"{API}/bookings/{booking['id']}", headers=auth(r1)).json()["payment_status"] == "unpaid"


def test_payment_details_and_invoice_pdf(client, state):
    sa, soa, r1, r2 = state["sa"], state["soa"], state["r1"], state["r2"]
    order = state["tamper"]["paid_order"]
    pid = order["payment_id"]

    d = client.get(f"{API}/payments/{pid}/details", headers=auth(sa))
    assert d.status_code == 200, d.text
    d = d.json()
    assert d["status"] == "success" and d["amount"] == 1400
    assert d["society"]["name"] and d["payer"]["name"] and d["payer"]["unit"]
    assert d["purpose"]["lines"][0]["qty"] == 2 and d["purpose"]["lines"][0]["rate"] == 700
    assert d["paid_to"]["name"] and d["gateway"]["order_id"] == order["gateway_order_id"]
    assert d["invoice_no"].startswith("INV-BK")

    # the payer and their society admin can open it, nobody else
    assert client.get(f"{API}/payments/{pid}/details", headers=auth(r1)).status_code == 200
    assert client.get(f"{API}/payments/{pid}/details", headers=auth(soa)).status_code == 200
    assert client.get(f"{API}/payments/{pid}/details", headers=auth(r2)).status_code == 404
    assert client.get(f"{API}/payments/{pid}/details", headers=auth(state["soa_b"])).status_code == 404

    pdf = client.get(f"{API}/payments/{pid}/invoice", headers=auth(sa))
    assert pdf.status_code == 200 and pdf.headers["content-type"] == "application/pdf"
    assert pdf.content.startswith(b"%PDF") and len(pdf.content) > 2500
    assert f'Invoice-{d["txn_id"]}.pdf' in pdf.headers["content-disposition"]
    assert client.get(f"{API}/payments/{pid}/invoice", headers=auth(r2)).status_code == 404

    # an unpaid / failed payment has no invoice
    deal = _deal(client, soa, price=50)
    b = client.post(f"{API}/bookings", headers=auth(r1), json={"deal_id": deal, "units_count": 1, "slot": "09:00-12:00"}).json()
    o = client.post(f"{API}/payments/initiate", headers=auth(r1), json={"ref_type": "booking", "ref_id": b["id"]}).json()
    assert client.get(f"{API}/payments/{o['payment_id']}/invoice", headers=auth(r1)).status_code == 400

    import os
    keep = os.environ.get("REPORT_SAMPLE_DIR")
    if keep:  # a hand-inspection hook
        open(os.path.join(keep, "invoice.pdf"), "wb").write(pdf.content)


def test_report_section_exports(client, state):
    import io

    from openpyxl import load_workbook

    sa = state["sa"]
    q = "date_from=2026-01-01&date_to=2026-12-31"
    for section in ("revenue", "societies", "bookings", "maintenance", "complaints", "residents"):
        pdf = client.get(f"{API}/reports/export?format=pdf&section={section}&{q}", headers=auth(sa))
        assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF"), section
        xl = client.get(f"{API}/reports/export?format=xlsx&section={section}&{q}", headers=auth(sa))
        assert xl.status_code == 200 and xl.content[:2] == b"PK", section
        assert f"-{section}-" in xl.headers["content-disposition"]
    sheets = load_workbook(io.BytesIO(client.get(f"{API}/reports/export?format=xlsx&section=bookings&{q}", headers=auth(sa)).content)).sheetnames
    assert sheets == ["Bookings & deals"]
    assert client.get(f"{API}/reports/export?format=pdf&section=nope&{q}", headers=auth(sa)).status_code == 422
    assert client.get(f"{API}/reports/export?format=pdf&section=revenue&{q}", headers=auth(state["soa"])).status_code == 403
