"""Excel / PDF of ONE chart of the Reports page (its own period, its own file)."""

import io
from datetime import date, datetime
from xml.sax.saxutils import escape

from app.modules.dashboard.report_export import to_xlsx

SECTIONS = {
    "revenue": ("Revenue over time", ["Revenue", "Payments"]),
    "societies": ("Revenue by society", ["Societies"]),
    "bookings": ("Bookings over time", ["Bookings & deals"]),
    "maintenance": ("Maintenance billed vs collected", ["Maintenance"]),
    "complaints": ("Complaints raised vs resolved", ["Complaints"]),
    "residents": ("Residents joined", ["Residents"]),
}


def to_section_xlsx(r: dict, section: str) -> bytes:
    """The full workbook reduced to the sheets that belong to one chart."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(to_xlsx(r)))
    keep = SECTIONS[section][1]
    for ws in list(wb.worksheets):
        if ws.title not in keep:
            wb.remove(ws)
    wb.active = 0
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_section_pdf(r: dict, section: str) -> bytes:
    """One page: Odeslo header, the chart's title and period, its chart, then its table."""
    from reportlab.graphics.charts.barcharts import HorizontalBarChart, VerticalBarChart
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from app.core import pdf_brand as brand

    title = SECTIONS[section][0]
    f, ts = r["filters"], r["timeseries"]
    page = landscape(A4)
    W = page[0] - 28 * mm

    def esc(v) -> str:
        return escape(str(v if v is not None else "-").encode("cp1252", "replace").decode("cp1252"))

    def rs(v: float) -> str:
        whole, frac = f"{abs(v):.2f}".split(".")
        if len(whole) > 3:
            head, tail, parts = whole[:-3], whole[-3:], []
            while len(head) > 2:
                parts.insert(0, head[-2:])
                head = head[:-2]
            if head:
                parts.insert(0, head)
            whole = ",".join(parts) + "," + tail
        return f"{'-' if v < 0 else ''}Rs. {whole}.{frac}"

    def st(n, **kw):
        return ParagraphStyle(n, fontName=kw.pop("font", "Helvetica"), textColor=kw.pop("color", brand.INK), **kw)

    body = st("b", fontSize=8.5, leading=11)
    numr = st("n", fontSize=8.5, leading=11, alignment=TA_RIGHT)
    th = st("th", fontSize=8, leading=10, color=colors.white, font="Helvetica-Bold")
    thr = st("thr", fontSize=8, leading=10, color=colors.white, font="Helvetica-Bold", alignment=TA_RIGHT)

    def table(headers, rows, widths, right=()):
        data = [[Paragraph(esc(h), thr if i in right else th) for i, h in enumerate(headers)]]
        data += [[Paragraph(esc(v), numr if i in right else body) for i, v in enumerate(row)] for row in rows]
        t = Table(data, colWidths=[W * w for w in widths], repeatRows=1)
        style = [("BACKGROUND", (0, 0), (-1, 0), brand.VIOLET), ("LINEBELOW", (0, 1), (-1, -1), 0.4, brand.LINE), ("TOPPADDING", (0, 0), (-1, -1), 4),
                 ("BOTTOMPADDING", (0, 0), (-1, -1), 4), ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6)]
        style += [("BACKGROUND", (0, i), (-1, i), brand.PANEL) for i in range(2, len(data), 2)]
        t.setStyle(TableStyle(style))
        return t

    pal = [brand.VIOLET, colors.HexColor("#2f9e44")]

    def bars(labels, series, note, money_axis=False):
        d = Drawing(W, 200)
        bc = VerticalBarChart()
        bc.x, bc.y, bc.width, bc.height = 50, 34, W - 70, 140
        bc.data = series
        bc.categoryAxis.categoryNames = labels
        bc.categoryAxis.labels.fontSize = 6.5
        bc.categoryAxis.labels.fontName = bc.valueAxis.labels.fontName = "Helvetica"
        bc.categoryAxis.labels.angle = 40 if len(labels) > 14 else 0
        bc.categoryAxis.labels.boxAnchor = "ne" if len(labels) > 14 else "n"
        bc.valueAxis.labels.fontSize = 7
        bc.valueAxis.valueMin = 0
        if money_axis:
            bc.valueAxis.labelTextFormat = lambda v: f"{v / 1000:.1f}k" if v >= 1000 else f"{v:,.0f}"
        for i in range(len(series)):
            bc.bars[i].fillColor = pal[i % len(pal)]
        bc.bars.strokeColor = None
        d.add(bc)
        d.add(String(50, 186, note, fontSize=8, fontName="Helvetica-Bold", fillColor=brand.MUTED))
        return d

    labels = [t["label"] for t in ts]
    if section == "revenue":
        chart = bars(labels, [[t["booking"] for t in ts] or [0], [t["maintenance"] for t in ts] or [0]], "Violet = group deals,  green = maintenance", True)
        tbl = table(["Period", "Group deals", "Maintenance", "Refunds", "Net", "Payments"],
                    [[t["label"], rs(t["booking"]), rs(t["maintenance"]), rs(t["refunds"]), rs(t["net"]), t["payments"]] for t in ts], [0.16, 0.2, 0.2, 0.16, 0.18, 0.1], (1, 2, 3, 4, 5))
    elif section == "societies":
        soc = r["revenue_by_society"]
        chart = Drawing(W, 40 + 16 * max(1, len(soc)))
        hb = HorizontalBarChart()
        hb.x, hb.y, hb.width, hb.height = 150, 10, W - 180, 14 * max(1, len(soc))
        hb.data = [[s["total"] for s in reversed(soc)] or [0]]
        hb.categoryAxis.categoryNames = [s["name"][:30] for s in reversed(soc)] or [""]
        hb.categoryAxis.labels.fontSize = 7
        hb.categoryAxis.labels.fontName = hb.valueAxis.labels.fontName = "Helvetica"
        hb.valueAxis.valueMin = 0
        hb.bars[0].fillColor = brand.VIOLET
        hb.bars.strokeColor = None
        chart.add(hb)
        tbl = table(["Society", "Group deals", "Maintenance", "Refunded", "Net revenue"],
                    [[s["society"], rs(s["booking_revenue"]), rs(s["maintenance_revenue"]), rs(s["refunded"]), rs(s["net_revenue"])] for s in r["societies"]], [0.3, 0.18, 0.18, 0.16, 0.18], (1, 2, 3, 4))
    elif section == "bookings":
        chart = bars(labels, [[t["bookings"] for t in ts] or [0], [t["units"] for t in ts] or [0]], "Violet = bookings,  green = units booked")
        tbl = table(["Period", "Bookings", "Units booked"], [[t["label"], t["bookings"], t["units"]] for t in ts], [0.5, 0.25, 0.25], (1, 2))
    elif section == "maintenance":
        m = r["maintenance_by_period"]
        chart = bars([x["period"] for x in m], [[x["billed"] for x in m] or [0], [x["collected"] for x in m] or [0]], "Violet = billed,  green = collected", True)
        tbl = table(["Period", "Bills", "Paid", "Overdue", "Billed", "Collected", "Collection"],
                    [[x["period"], x["bills"], x["paid"], x["overdue"], rs(x["billed"]), rs(x["collected"]), f"{x['rate']}%"] for x in m], [0.14, 0.1, 0.1, 0.1, 0.2, 0.2, 0.16], (1, 2, 3, 4, 5, 6))
    elif section == "complaints":
        chart = bars(labels, [[t["complaints"] for t in ts] or [0], [t["resolved"] for t in ts] or [0]], "Violet = raised,  green = resolved")
        tbl = table(["Period", "Raised", "Resolved"], [[t["label"], t["complaints"], t["resolved"]] for t in ts], [0.5, 0.25, 0.25], (1, 2))
    else:
        chart = bars(labels, [[t["residents"] for t in ts] or [0]], "Residents joined")
        tbl = table(["Period", "Residents joined"], [[t["label"], t["residents"]] for t in ts], [0.6, 0.4], (1,))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=13 * mm, bottomMargin=16 * mm, title=f"Odeslo - {title}", author="Odeslo")
    brand_cell = Table([[brand.logo(13), Paragraph('<font name="Helvetica-Bold" size="20" color="#5446d9">Odeslo</font><br/><font size="8" color="#6e6e78">Platform report</font>',
                                                    st("hd", fontSize=20, leading=22))]], colWidths=[17 * mm, 70 * mm])
    brand_cell.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    when = f"{date.fromisoformat(f['date_from']):%d %b %Y} to {date.fromisoformat(f['date_to']):%d %b %Y}"
    right_text = (f'<font name="Helvetica-Bold" size="12">{esc(title)}</font><br/>'
                  f'<font size="8.5" color="#6e6e78">{when} | {esc(", ".join(f["societies"]))} | grouped {esc(f["group_by"])}ly</font>')
    head = Table([[brand_cell, Paragraph(right_text, st("hr", fontSize=10, leading=14, alignment=TA_RIGHT))]], colWidths=[W * 0.4, W * 0.6])
    head.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    story = [head, Spacer(1, 3 * mm), HRFlowable(width="100%", thickness=2.2, color=brand.VIOLET, spaceAfter=4 * mm), chart, Spacer(1, 4 * mm), tbl]

    def footer(canvas, doc_):
        brand.watermark(canvas, page, size_mm=120)
        canvas.saveState()
        canvas.setStrokeColor(brand.LINE)
        canvas.line(14 * mm, 11 * mm, page[0] - 14 * mm, 11 * mm)
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(brand.MUTED)
        canvas.drawString(14 * mm, 7 * mm, f"Odeslo - {title}. Generated {datetime.fromisoformat(r['generated_at']):%d %b %Y, %I:%M %p}")
        canvas.drawRightString(page[0] - 14 * mm, 7 * mm, f"Page {doc_.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()
