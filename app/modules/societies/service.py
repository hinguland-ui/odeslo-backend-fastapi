import csv
import io
import re

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.enums import DEFAULT_FEATURES, OccupancyStatus, ResidentStatus, SocietyStatus, UnitType
from app.core.errors import bad_request, conflict, forbidden, not_found
from app.modules.residents.models import Resident
from app.modules.societies.models import Society, Unit, Wing


def generate_society_code(db: Session, name: str) -> str:
    """Short unique code such as AJR001 for 'Ajit Residency'."""
    words = [w for w in re.findall(r"[A-Za-z]+", name.upper())]
    if len(words) >= 2:
        prefix = words[0][:2] + words[1][:1]
    elif words:
        prefix = words[0][:3]
    else:
        prefix = "SOC"
    prefix = prefix.ljust(3, "X")
    n = 1
    while db.scalar(select(Society.id).where(Society.code == f"{prefix}{n:03d}")):
        n += 1
    return f"{prefix}{n:03d}"


def get_society_or_404(db: Session, society_id: int) -> Society:
    society = db.get(Society, society_id)
    if not society:
        raise not_found("Society")
    return society


def is_feature_enabled(society: Society, feature: str) -> bool:
    return {**DEFAULT_FEATURES, **(society.features or {})}.get(feature, False)


def ensure_feature_enabled(db: Session, society_id: int | None, feature: str) -> None:
    if society_id is None:
        return
    society = get_society_or_404(db, society_id)
    if not is_feature_enabled(society, feature):
        raise forbidden(f"The '{feature}' feature is not enabled for this society")


def add_wing(db: Session, society_id: int, name: str) -> Wing:
    name = name.strip().upper() if len(name.strip()) <= 2 else name.strip()
    if db.scalar(select(Wing).where(Wing.society_id == society_id, Wing.name == name)):
        raise conflict(f"Wing '{name}' already exists")
    wing = Wing(society_id=society_id, name=name)
    db.add(wing)
    db.flush()
    return wing


def add_unit(db: Session, wing: Wing, unit_no: str, unit_type: str, floor: int | None) -> Unit:
    unit_no = unit_no.strip()
    if db.scalar(select(Unit).where(Unit.wing_id == wing.id, Unit.unit_no == unit_no)):
        raise conflict(f"Unit {wing.name}-{unit_no} already exists")
    unit = Unit(society_id=wing.society_id, wing_id=wing.id, unit_no=unit_no, unit_type=unit_type, floor=floor)
    db.add(unit)
    db.flush()
    return unit


def society_stats(db: Session, society_id: int) -> dict:
    def count(stmt):
        return db.scalar(stmt) or 0

    return {
        "wings": count(select(func.count()).select_from(Wing).where(Wing.society_id == society_id)),
        "units": count(select(func.count()).select_from(Unit).where(Unit.society_id == society_id)),
        "occupied_units": count(select(func.count()).select_from(Unit).where(
            Unit.society_id == society_id, Unit.occupancy_status != OccupancyStatus.VACANT)),
        "active_residents": count(select(func.count()).select_from(Resident).where(
            Resident.society_id == society_id, Resident.status == ResidentStatus.ACTIVE)),
        "pending_residents": count(select(func.count()).select_from(Resident).where(
            Resident.society_id == society_id, Resident.status == ResidentStatus.PENDING_APPROVAL)),
    }


def refresh_unit_occupancy(db: Session, unit_id: int) -> None:
    """Occupancy = tenant if an active tenant lives there, else owner if an active owner, else vacant."""
    unit = db.get(Unit, unit_id)
    if not unit:
        return
    types = set(db.scalars(select(Resident.resident_type).where(
        Resident.unit_id == unit_id, Resident.status == ResidentStatus.ACTIVE)).all())
    if "tenant" in types:
        unit.occupancy_status = OccupancyStatus.TENANT
    elif "owner" in types:
        unit.occupancy_status = OccupancyStatus.OWNER
    else:
        unit.occupancy_status = OccupancyStatus.VACANT


def _parse_rows(filename: str, content: bytes) -> list[dict]:
    if filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(content), read_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h or "").strip().lower() for h in rows[0]]
        return [dict(zip(headers, r)) for r in rows[1:] if any(c is not None for c in r)]
    text = content.decode("utf-8-sig")
    return [{(k or "").strip().lower(): v for k, v in row.items()} for row in csv.DictReader(io.StringIO(text))]


def bulk_upload_units(db: Session, society_id: int, filename: str, content: bytes) -> dict:
    """Columns: wing, unit_no, unit_type (flat/duplex/shop/row_house), floor."""
    try:
        rows = _parse_rows(filename, content)
    except Exception as exc:
        raise bad_request(f"Could not read file: {exc}")
    created, skipped, errors = 0, 0, []
    wings = {w.name: w for w in db.scalars(select(Wing).where(Wing.society_id == society_id)).all()}
    valid_types = {t.value for t in UnitType}
    for i, row in enumerate(rows, start=2):
        wing_name = str(row.get("wing") or "").strip()
        unit_no = str(row.get("unit_no") or row.get("unit") or "").strip()
        if unit_no.endswith(".0"):
            unit_no = unit_no[:-2]
        unit_type = str(row.get("unit_type") or "flat").strip().lower()
        floor_raw = row.get("floor")
        if not wing_name or not unit_no:
            errors.append(f"Row {i}: wing and unit_no are required")
            continue
        if unit_type not in valid_types:
            errors.append(f"Row {i}: invalid unit_type '{unit_type}'")
            continue
        try:
            floor = int(float(floor_raw)) if floor_raw not in (None, "") else None
        except ValueError:
            errors.append(f"Row {i}: invalid floor '{floor_raw}'")
            continue
        wing = wings.get(wing_name) or wings.get(wing_name.upper())
        if not wing:
            wing = add_wing(db, society_id, wing_name)
            wings[wing.name] = wing
        if db.scalar(select(Unit.id).where(Unit.wing_id == wing.id, Unit.unit_no == unit_no)):
            skipped += 1
            continue
        add_unit(db, wing, unit_no, unit_type, floor)
        created += 1
    return {"created": created, "skipped": skipped, "errors": errors}


def ensure_society_active(society: Society) -> None:
    if society.status != SocietyStatus.ACTIVE:
        raise bad_request("This society is not active yet")
